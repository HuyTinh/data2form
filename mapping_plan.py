"""Validation and workbook loading for versioned parent/detail mapping plans."""
from __future__ import annotations

from copy import deepcopy
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from form_fields import validate_mappings


PLAN_VERSION = 1


def _require_text(value: Any, label: str) -> str:
    text = str(value or "").strip()
    if not text:
        raise ValueError(f"{label} is required")
    return text


def _validate_mappings(mappings: Any, columns: list[str], label: str) -> None:
    if not isinstance(mappings, dict):
        raise ValueError(f"{label} mappings must be an object")
    if not mappings:
        raise ValueError(f"{label} must have at least one field mapping")
    for source_column, mapping in mappings.items():
        if not isinstance(mapping, dict):
            raise ValueError(f"{label} mapping '{source_column}' must be an object")
        _require_text(mapping.get("selector"), f"{label} selector for '{source_column}'")
    errors = validate_mappings(mappings, columns)
    if errors:
        raise ValueError(f"{label}: " + "; ".join(errors))


def _validate_raw_headers(excel_path: str, sheet_names: set[str]) -> None:
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        for sheet_name in sheet_names:
            worksheet = workbook[sheet_name]
            raw_headers = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            headers = [str(value).strip() if value is not None else "" for value in raw_headers]
            if any(not header for header in headers):
                raise ValueError(f"worksheet '{sheet_name}' contains a blank header")
            seen: set[str] = set()
            duplicates: set[str] = set()
            for header in headers:
                if header in seen:
                    duplicates.add(header)
                seen.add(header)
            if duplicates:
                raise ValueError(f"worksheet '{sheet_name}' contains duplicate header(s): " + ", ".join(sorted(duplicates)))
    finally:
        workbook.close()


def load_mapping_plan_workbook(excel_path: str, mapping_plan: dict[str, Any]) -> dict[str, Any]:
    """Validate a complete plan and return parent rows plus detail rows grouped by parent key.

    The function is read-only: it opens only sheets referenced by the plan and makes
    no browser or database mutations. Parent/detail key values are trimmed strings.
    """
    if not isinstance(mapping_plan, dict) or type(mapping_plan.get("version")) is not int or mapping_plan.get("version") != PLAN_VERSION:
        raise ValueError(f"mapping_plan version must be {PLAN_VERSION}")

    normalized_plan = deepcopy(mapping_plan)
    parent = normalized_plan.get("parent")
    if not isinstance(parent, dict):
        raise ValueError("mapping_plan parent must be an object")
    parent_sheet = _require_text(parent.get("sheet_name"), "parent sheet_name")
    parent_key_column = _require_text(parent.get("key_column"), "parent key_column")
    parent["sheet_name"] = parent_sheet
    parent["key_column"] = parent_key_column
    submit_selector = _require_text(normalized_plan.get("submit_selector"), "submit_selector")
    normalized_plan["submit_selector"] = submit_selector
    detail_tables = normalized_plan.get("detail_tables")
    if not isinstance(detail_tables, list) or not detail_tables:
        raise ValueError("mapping_plan must contain at least one detail table")

    table_ids: set[str] = set()
    table_sheets: set[str] = {parent_sheet}
    for index, table in enumerate(detail_tables, start=1):
        if not isinstance(table, dict):
            raise ValueError(f"detail_tables[{index - 1}] must be an object")
        table_id = _require_text(table.get("id"), f"detail_tables[{index - 1}].id")
        if table_id in table_ids:
            raise ValueError(f"duplicate detail table id '{table_id}'")
        table_ids.add(table_id)
        table_name = _require_text(table.get("name"), f"detail table '{table_id}' name")
        sheet_name = _require_text(table.get("sheet_name"), f"detail table '{table_id}' sheet_name")
        if sheet_name in table_sheets:
            raise ValueError(f"detail table sheet '{sheet_name}' must be distinct")
        table_sheets.add(sheet_name)
        key_column = _require_text(table.get("parent_key_column"), f"detail table '{table_id}' parent_key_column")
        try:
            row_save_timeout = int(table.get("row_save_timeout", 8000))
        except (TypeError, ValueError) as exc:
            raise ValueError(f"detail table '{table_id}' row_save_timeout must be an integer") from exc
        if not 1000 <= row_save_timeout <= 60000:
            raise ValueError(f"detail table '{table_id}' row_save_timeout must be between 1000 and 60000 ms")
        table.update({
            "id": table_id,
            "name": table_name,
            "sheet_name": sheet_name,
            "parent_key_column": key_column,
            "row_save_timeout": row_save_timeout,
        })

    try:
        with pd.ExcelFile(excel_path) as workbook:
            available_sheets = set(workbook.sheet_names)
    except Exception as exc:
        raise ValueError(f"Unable to open Excel workbook: {exc}") from exc

    missing_sheets = table_sheets - available_sheets
    if missing_sheets:
        raise ValueError("unknown sheet(s): " + ", ".join(sorted(missing_sheets)))
    _validate_raw_headers(excel_path, table_sheets)

    parent_df = pd.read_excel(excel_path, sheet_name=parent_sheet, dtype=str).fillna("")
    if parent_df.empty:
        raise ValueError(f"parent sheet '{parent_sheet}' contains no parent records")
    parent_df = parent_df.loc[parent_df.astype(str).apply(lambda row: row.str.strip().ne("").any(), axis=1)].reset_index(drop=True)
    if parent_df.empty:
        raise ValueError(f"parent sheet '{parent_sheet}' contains no parent records")
    if parent_key_column not in parent_df.columns:
        raise ValueError(f"parent sheet '{parent_sheet}' is missing key column '{parent_key_column}'")
    _validate_mappings(parent.get("mappings"), list(parent_df.columns), "parent")

    parent_keys = parent_df[parent_key_column].astype(str).str.strip()
    if (parent_keys == "").any():
        raise ValueError(f"parent sheet '{parent_sheet}' contains a blank parent key")
    duplicate_keys = parent_keys[parent_keys.duplicated()].unique().tolist()
    if duplicate_keys:
        raise ValueError("duplicate parent key(s): " + ", ".join(duplicate_keys))

    parent_rows = parent_df.to_dict(orient="records")
    for row, key in zip(parent_rows, parent_keys.tolist()):
        row[parent_key_column] = key
    valid_parent_keys = set(parent_keys.tolist())
    detail_rows: dict[str, dict[str, list[dict[str, Any]]]] = {}

    for table in detail_tables:
        table_id = table["id"].strip()
        sheet_name = table["sheet_name"].strip()
        key_column = table["parent_key_column"].strip()
        detail_df = pd.read_excel(excel_path, sheet_name=sheet_name, dtype=str).fillna("")
        detail_df = detail_df.loc[detail_df.astype(str).apply(lambda row: row.str.strip().ne("").any(), axis=1)].reset_index(drop=True)
        if key_column not in detail_df.columns:
            raise ValueError(f"detail table '{table_id}' sheet '{sheet_name}' is missing key column '{key_column}'")
        _validate_mappings(table.get("mappings"), list(detail_df.columns), f"detail table '{table_id}'")

        grouped: dict[str, list[dict[str, Any]]] = {key: [] for key in valid_parent_keys}
        for source_row in detail_df.to_dict(orient="records"):
            if not any(str(value or "").strip() for value in source_row.values()):
                continue
            key = str(source_row.get(key_column, "")).strip()
            if not key:
                raise ValueError(f"detail table '{table_id}' contains a blank parent key")
            if key not in valid_parent_keys:
                raise ValueError(f"detail table '{table_id}' parent key '{key}' has no matching parent")
            source_row[key_column] = key
            grouped[key].append(source_row)
        detail_rows[table_id] = grouped

    return {
        "plan": normalized_plan,
        "parent_rows": parent_rows,
        "detail_rows": detail_rows,
    }
